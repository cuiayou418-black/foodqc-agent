"""品控日报智能体 · 多格式导出（Excel / PDF）

- Excel（openpyxl）：概览 / 合格率 / 缺陷分布 / 异常批次 / 趋势 / 质检明细 / 仓储记录 /
  溯源审计 共 8 个 Sheet。质检明细带标准号与限值出处 —— 每个数字都能翻回国标原文；
- PDF（reportlab）：A4 排版，自动注册中文字体（优先微软雅黑，回退 STSong-Light CID 字体）。
"""
from __future__ import annotations

import os

from .models import CollectedData, DailyMetrics

# ---------------------------------------------------------------- 公共样式常量

RED = "B81C22"
LIGHT_RED = "FEF0EF"
LIGHT_AMBER = "FDF6EC"
GREY_BG = "F2F3F5"


def _ensure_parent(path: str) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)


# ---------------------------------------------------------------- Excel

def export_excel(m: DailyMetrics, data: CollectedData, meta: dict, path: str) -> str:
    """导出 Excel 日报。返回文件路径。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    _ensure_parent(path)
    wb = Workbook()

    thin = Side(style="thin", color="DEE0E3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=RED)
    head_font = Font(color="FFFFFF", bold=True, size=11)
    fail_fill = PatternFill("solid", fgColor=LIGHT_RED)
    warn_fill = PatternFill("solid", fgColor=LIGHT_AMBER)
    title_font = Font(bold=True, size=14, color=RED)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def sheet(name: str, title: str, widths: list):
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=1, value=title).font = title_font
        ws.cell(row=2, column=1,
                value=f"报告日期 {m.report_date} ｜ {m.plant} ｜ 生成时间 {m.generated_at} ｜ "
                      f"RunID {meta.get('run_id', '')}").font = Font(size=9, color="8F959E")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A4"
        return ws

    def header_row(ws, row: int, headers: list):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.fill, cell.font, cell.border, cell.alignment = head_fill, head_font, border, center

    def data_row(ws, row: int, values: list, fill=None, num_cols=()):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.border = border
            cell.alignment = Alignment(horizontal="right" if c in num_cols else "left",
                                       vertical="center", wrap_text=True)
            if c in num_cols and isinstance(v, float):
                cell.number_format = "0.00%" if abs(v) <= 1.5 else "0.00"
            if fill:
                cell.fill = fill

    pct = lambda x: f"{x * 100:.2f}%"

    # ---- 1. 概览 ----
    ws = sheet("概览", f"{meta.get('title', '食用油中控品控日报')} · 核心指标", [22, 16, 60])
    header_row(ws, 4, ["指标", "数值", "说明"])
    rows = [
        ("总体合格率", pct(m.pass_rate), f"{m.total_tests - m.fail_tests}/{m.total_tests} 项次合格"),
        ("一次交验合格率", pct(m.batch_pass_rate), "全部项目一次通过的批次占比"),
        ("生产批次", m.batch_count, "当日投入生产并完成检测的批次数"),
        ("计划产量(吨)", m.planned_qty, ""), ("实际产量(吨)", m.actual_qty, ""),
        ("产量达成率", pct(m.achievement), ""),
        ("不合格项次", m.fail_tests, ""),
        ("异常批次", len(m.abnormal_batches),
         f"不合格 {sum(1 for b in m.abnormal_batches if b.level == 'critical')} · "
         f"预警 {sum(1 for b in m.abnormal_batches if b.level == 'warning')} · "
         f"提示 {sum(1 for b in m.abnormal_batches if b.level == 'info')}"),
        ("接近限值项次", m.near_limit_count, "限值利用率 ≥ 80%"),
        ("仓储环境超限", len(m.warehouse_violations), "库温/湿度超出配置范围"),
        ("预警条数", len(m.alerts),
         f"严重 {sum(1 for a in m.alerts if a.level == 'critical')} · "
         f"警告 {sum(1 for a in m.alerts if a.level == 'warning')} · "
         f"提示 {sum(1 for a in m.alerts if a.level == 'info')}"),
    ]
    r = 5
    for name, val, note in rows:
        fill = fail_fill if (name == "总体合格率" and m.pass_rate < 0.95) else None
        data_row(ws, r, [name, val, note], fill=fill, num_cols=(2,) if isinstance(val, (int, float)) else ())
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="当日预警").font = Font(bold=True)
    r += 1
    for a in m.alerts:
        data_row(ws, r, [a.level.upper(), a.title, a.detail], num_cols=())
        r += 1

    # ---- 2. 合格率 ----
    ws = sheet("合格率", "合格率分析（分产品 / 分检测项目）", [18, 10, 12, 12, 12, 14, 14, 14])
    header_row(ws, 4, ["产品", "批次", "检测项次", "不合格项次", "项目合格率",
                       "一次交验合格率", "计划产量(t)", "实际产量(t)"])
    r = 5
    for p in m.products:
        data_row(ws, r, [p.product, p.batches, p.total_tests, p.fail_tests, p.pass_rate,
                         p.batch_pass_rate, p.planned_qty, p.actual_qty],
                 fill=fail_fill if p.fail_tests else None, num_cols=(5, 6))
        r += 1
    data_row(ws, r, ["合计", m.batch_count, m.total_tests, m.fail_tests, m.pass_rate,
                     m.batch_pass_rate, m.planned_qty, m.actual_qty], num_cols=(5, 6))
    r += 3
    ws.cell(row=r - 1, column=1, value="分检测项目").font = Font(bold=True)
    header_row(ws, r, ["检测项目", "单位", "检测项次", "不合格项次", "合格率", "涉及批次", "", ""])
    r += 1
    for it in m.items:
        data_row(ws, r, [it.item, it.unit, it.total, it.fail, it.pass_rate,
                         ", ".join(it.fail_batches), "", ""],
                 fill=fail_fill if it.fail else None, num_cols=(5,))
        r += 1

    # ---- 3. 缺陷分布 ----
    ws = sheet("缺陷分布", "缺陷分布（按检测项目）", [18, 10, 14, 14, 46])
    header_row(ws, 4, ["检测项目", "单位", "不合格项次", "占全部缺陷比例", "涉及批次"])
    r = 5
    for d in m.defect_distribution:
        data_row(ws, r, [d["item"], d["unit"], d["fail"], d["share"], ", ".join(d["fail_batches"])],
                 fill=fail_fill, num_cols=(4,))
        r += 1
    if not m.defect_distribution:
        data_row(ws, r, ["当日无不合格项次", "", "", "", ""])

    # ---- 4. 异常批次 ----
    ws = sheet("异常批次", "异常批次明细", [16, 18, 16, 10, 12, 70])
    header_row(ws, 4, ["批次号", "产品", "产线", "级别", "最大利用率", "异常原因"])
    r = 5
    for b in m.abnormal_batches:
        fill = fail_fill if b.level == "critical" else (warn_fill if b.level == "warning" else None)
        data_row(ws, r, [b.batch_id, b.product, b.line,
                         {"critical": "不合格", "warning": "预警", "info": "提示"}[b.level],
                         b.worst_ratio, "；".join(b.reasons)], fill=fill, num_cols=(5,))
        r += 1
    if not m.abnormal_batches:
        data_row(ws, r, ["当日无异常批次"] + [""] * 5)

    # ---- 5. 趋势 ----
    ws = sheet("趋势", f"近 {len(m.trend_points)} 日趋势分析", [14, 12, 12, 12, 12, 40])
    header_row(ws, 4, ["日期", "合格率", "检测项次", "不合格项次", "异常批次", ""])
    r = 5
    for tp in m.trend_points:
        data_row(ws, r, [tp.date, tp.pass_rate, tp.total, tp.fail, tp.abnormal_batches, ""],
                 fill=fail_fill if tp.fail else None, num_cols=(2,))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=m.trend_summary.get("comment", "")).font = Font(bold=True)
    r += 2
    ws.cell(row=r - 1, column=1, value="分检测项目均值").font = Font(bold=True)
    item_names = sorted({t.item for t in m.item_trend})
    header_row(ws, r, ["日期"] + item_names)
    r += 1
    lookup = {(t.date, t.item): t for t in m.item_trend}
    for d in sorted({t.date for t in m.item_trend}):
        cells = [d]
        for n in item_names:
            t = lookup.get((d, n))
            cells.append(f"{t.mean_value:g}{t.unit}" if t else "-")
        data_row(ws, r, cells)
        r += 1

    # ---- 6. 质检明细（溯源核心） ----
    ws = sheet("质检明细", "质检记录明细（含标准出处，可逐条回溯）",
               [22, 16, 16, 14, 10, 8, 10, 16, 30, 8, 18, 12, 12, 14, 20])
    header_row(ws, 4, ["记录号", "批次号", "产品", "检测项目", "实测值", "单位", "方向",
                       "限值", "依据标准", "判定", "限值出处", "检测时间", "检验员", "利用率", "来源系统"])
    r = 5
    for rec in data.quality:
        data_row(ws, r, [rec.record_id, rec.batch_id, rec.product, rec.item, rec.value, rec.unit,
                         rec.direction, rec.limit_value, rec.standard, rec.verdict,
                         rec.limit_source, rec.tested_at, rec.inspector, round(rec.ratio, 4), rec.source],
                 fill=fail_fill if rec.verdict == "不合格" else None, num_cols=(5, 8, 14))
        r += 1

    # ---- 7. 仓储记录 ----
    ws = sheet("仓储记录", "仓储流水明细", [22, 16, 14, 8, 10, 10, 10, 18, 20])
    header_row(ws, 4, ["记录号", "批次号", "仓库", "事件", "数量(t)", "库温(℃)", "湿度(%)", "发生时间", "来源系统"])
    r = 5
    limits = meta.get("warehouse_limits", {})
    for w in data.warehouse:
        viol = ((w.temp is not None and not (limits.get("temp_min", -999) <= w.temp <= limits.get("temp_max", 999)))
                or (w.humidity is not None and not (limits.get("humidity_min", -999) <= w.humidity <= limits.get("humidity_max", 999))))
        data_row(ws, r, [w.record_id, w.batch_id, w.warehouse, w.event, w.qty, w.temp, w.humidity,
                         w.occurred_at, w.source],
                 fill=warn_fill if viol else None, num_cols=(5,))
        r += 1

    # ---- 8. 溯源与审计 ----
    ws = sheet("溯源与审计", "数据溯源与审计信息", [14, 12, 10, 40, 40, 20])
    header_row(ws, 4, ["来源系统", "类型", "记录数", "数据窗口", "记录指纹(SHA-256 前16位)", "抽取时间"])
    r = 5
    for p in meta.get("provenance", []):
        data_row(ws, r, [p["system"], p["source_type"], p["record_count"],
                         f"{p['window_start']} ~ {p['window_end']}", p["checksum"][:16], p["extracted_at"]],
                 num_cols=(3,))
        r += 1
    r += 1
    for label, val in [("原始数据快照指纹", meta.get("data_checksum", "")),
                       ("本报告指纹", meta.get("fingerprint", "")),
                       ("审计日志路径", meta.get("audit_path", "")),
                       ("采集快照目录", meta.get("snapshot_dir", ""))]:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=val).alignment = left
        r += 1

    del wb["Sheet"]  # 删除默认空表
    wb.save(path)
    return path


# ---------------------------------------------------------------- PDF

def _register_cn_font():
    """注册中文字体：优先系统 TTF，回退 reportlab 内置 CID 字体（无需字体文件）。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        ("msyh", r"C:\Windows\Fonts\msyh.ttc"),
        ("msyh", r"C:\Windows\Fonts\msyh.ttf"),
        ("simhei", r"C:\Windows\Fonts\simhei.ttf"),
        ("simsun", r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for name, path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                return name
            except Exception:
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def export_pdf(m: DailyMetrics, meta: dict, path: str) -> str:
    """导出 PDF 日报（A4，中文字体，页脚含报告指纹与页码）。返回文件路径。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.platypus import (HRFlowable, KeepTogether, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    _ensure_parent(path)
    font = _register_cn_font()

    RED = colors.HexColor("#B81C22")
    GREY = colors.HexColor("#646A73")
    LIGHT = colors.HexColor("#F2F3F5")
    FAIL_BG = colors.HexColor("#FEF0EF")
    WARN_BG = colors.HexColor("#FDF6EC")

    st_title = ParagraphStyle("title", fontName=font, fontSize=17, leading=22, textColor=RED)
    st_sub = ParagraphStyle("sub", fontName=font, fontSize=8.5, leading=12, textColor=GREY)
    st_h2 = ParagraphStyle("h2", fontName=font, fontSize=12.5, leading=16, textColor=RED,
                           spaceBefore=14, spaceAfter=6)
    st_h3 = ParagraphStyle("h3", fontName=font, fontSize=10, leading=13, textColor=colors.HexColor("#383A42"),
                           spaceBefore=8, spaceAfter=4)
    st_body = ParagraphStyle("body", fontName=font, fontSize=9, leading=13)
    st_cell = ParagraphStyle("cell", fontName=font, fontSize=8, leading=11)
    st_cell_b = ParagraphStyle("cellb", fontName=font, fontSize=8, leading=11, textColor=RED)
    st_alert = ParagraphStyle("alert", fontName=font, fontSize=9, leading=13)

    def table_style(header_bg=RED, extra=None):
        base = [
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), header_bg),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DEE0E3")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]
        return TableStyle(base + (extra or []))

    story = []

    # ---------- 报告头 ----------
    story.append(Paragraph(meta.get("title", "食用油中控品控日报"), st_title))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        f"报告日期：{m.report_date}　|　生产基地：{m.plant}　|　生成时间：{m.generated_at}<br/>"
        f"数据来源：{meta.get('source_type', '')}　|　运行编号：{meta.get('run_id', '')}", st_sub))
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=2, color=RED))

    # ---------- 校验提示 ----------
    errors = [v for v in m.validation_issues if v.level == "error"]
    if errors:
        story.append(Spacer(1, 6))
        story.append(Paragraph(f"<b>⛔ 数据校验未通过（{len(errors)} 项错误）</b>："
                               + "；".join(e.message for e in errors[:5])
                               + "。本报告禁止直接用于放行决策，请人工复核。", st_alert))

    # ---------- KPI ----------
    kpi = [["总体合格率", "一次交验合格率", "产量达成率", "异常批次", "当日预警"],
           [f"{m.pass_rate:.2%}", f"{m.batch_pass_rate:.2%}", f"{m.achievement:.1%}",
            f"{len(m.abnormal_batches)} 个", f"{len(m.alerts)} 条"]]
    story.append(Spacer(1, 10))
    story.append(Table(kpi, colWidths=[36 * mm] * 5,
                       style=TableStyle([
                           ("FONTNAME", (0, 0), (-1, -1), font),
                           ("FONTSIZE", (0, 0), (-1, 0), 7.5),
                           ("FONTSIZE", (0, 1), (-1, 1), 13),
                           ("BACKGROUND", (0, 0), (-1, 0), LIGHT),
                           ("TEXTCOLOR", (0, 1), (-1, 1), RED),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DEE0E3")),
                           ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                           ("TOPPADDING", (0, 0), (-1, -1), 5),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                       ])))

    # ---------- 一、合格率 ----------
    story.append(Paragraph("一、合格率分析", st_h2))
    rows = [["产品", "批次", "检测项次", "不合格项次", "项目合格率", "一次交验", "计划(t)", "实际(t)"]]
    for p in m.products:
        rows.append([p.product, p.batches, p.total_tests, p.fail_tests,
                     f"{p.pass_rate:.2%}", f"{p.batch_pass_rate:.2%}",
                     f"{p.planned_qty:g}", f"{p.actual_qty:g}"])
    rows.append(["合计", m.batch_count, m.total_tests, m.fail_tests,
                 f"{m.pass_rate:.2%}", f"{m.batch_pass_rate:.2%}",
                 f"{m.planned_qty:g}", f"{m.actual_qty:g}"])
    story.append(Table(rows, colWidths=[42 * mm, 12 * mm, 17 * mm, 20 * mm, 20 * mm, 18 * mm, 16 * mm, 16 * mm],
                       style=table_style(extra=[("FONTNAME", (0, -1), (-1, -1), font),
                                                ("BACKGROUND", (0, -1), (-1, -1), LIGHT)])))
    story.append(Paragraph("分检测项目", st_h3))
    rows = [["检测项目", "单位", "检测项次", "不合格项次", "合格率"]]
    for it in m.items:
        rows.append([it.item, it.unit, it.total, it.fail, f"{it.pass_rate:.2%}"])
    story.append(Table(rows, colWidths=[50 * mm, 20 * mm, 22 * mm, 24 * mm, 24 * mm], style=table_style()))

    # ---------- 二、缺陷分布 ----------
    story.append(Paragraph("二、缺陷分布", st_h2))
    if m.defect_distribution:
        rows = [["检测项目", "不合格项次", "占全部缺陷比例", "涉及批次"]]
        for d in m.defect_distribution:
            rows.append([d["item"], d["fail"], f"{d['share'] * 100:.1f}%",
                         Paragraph(", ".join(d["fail_batches"]), st_cell)])
        story.append(Table(rows, colWidths=[36 * mm, 22 * mm, 28 * mm, 74 * mm], style=table_style()))
    else:
        story.append(Paragraph("当日无不合格项次，缺陷分布为空。", st_body))

    # ---------- 三、异常批次 ----------
    story.append(Paragraph("三、异常批次明细", st_h2))
    if m.abnormal_batches:
        rows = [["批次号", "产品", "级别", "最大利用率", "异常原因"]]
        for b in m.abnormal_batches:
            rows.append([Paragraph(b.batch_id, st_cell), b.product,
                         {"critical": "不合格", "warning": "预警", "info": "提示"}[b.level],
                         f"{b.worst_ratio:.0%}", Paragraph("<br/>".join(b.reasons), st_cell)])
        t = Table(rows, colWidths=[26 * mm, 26 * mm, 14 * mm, 16 * mm, 78 * mm], style=table_style())
        story.append(t)
    else:
        story.append(Paragraph("当日无异常批次。", st_body))

    # ---------- 四、趋势分析 ----------
    story.append(Paragraph("四、趋势分析", st_h2))
    if len(m.trend_points) >= 2:
        story.append(Paragraph(m.trend_summary.get("comment", ""), st_body))
        rows = [["日期", "合格率", "检测项次", "不合格项次", "异常批次"]]
        for tp in m.trend_points:
            rows.append([tp.date, f"{tp.pass_rate:.2%}", tp.total, tp.fail, tp.abnormal_batches])
        story.append(Table(rows, colWidths=[36 * mm, 28 * mm, 26 * mm, 28 * mm, 26 * mm], style=table_style()))
        item_names = sorted({t.item for t in m.item_trend})
        if item_names:
            story.append(Paragraph("分检测项目均值", st_h3))
            rows = [["日期"] + item_names]
            lookup = {(t.date, t.item): t for t in m.item_trend}
            for d in sorted({t.date for t in m.item_trend}):
                cells = [d]
                for n in item_names:
                    t = lookup.get((d, n))
                    cells.append(f"{t.mean_value:g}{t.unit}" if t else "-")
                rows.append(cells)
            story.append(Table(rows, colWidths=[28 * mm] + [(160 - 28) / max(len(item_names), 1) * mm] * len(item_names),
                               style=table_style()))
    else:
        story.append(Paragraph("历史数据不足，连续运行数日后将自动生成近 7 日走势。", st_body))

    # ---------- 五、预警汇总 ----------
    story.append(Paragraph("五、预警汇总", st_h2))
    if m.alerts:
        for a in m.alerts:
            body = f"<b>[{a.level.upper()}] {a.title}</b><br/>{a.detail}"
            if a.refs:
                body += f"<br/><font color='#8F959E' size='7'>关联记录：{', '.join(a.refs[:12])}</font>"
            story.append(KeepTogether(Paragraph(body, st_alert)))
            story.append(Spacer(1, 3))
    else:
        story.append(Paragraph("当日无预警命中，各项指标处于受控范围。", st_body))

    # ---------- 六、数据溯源与审计 ----------
    story.append(Paragraph("六、数据溯源与审计", st_h2))
    rows = [["来源系统", "类型", "记录数", "数据窗口", "记录指纹(SHA-256 前16位)"]]
    for p in meta.get("provenance", []):
        rows.append([p["system"], p["source_type"], p["record_count"],
                     f"{p['window_start']} ~ {p['window_end']}", p["checksum"][:16]])
    story.append(Table(rows, colWidths=[22 * mm, 20 * mm, 14 * mm, 56 * mm, 38 * mm], style=table_style()))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        f"原始数据快照指纹：<font face='Courier' size='7.5'>{meta.get('data_checksum', '')[:32]}</font><br/>"
        f"本报告指纹：<font face='Courier' size='7.5'>{meta.get('fingerprint', '')}</font><br/>"
        f"审计日志：{meta.get('audit_path', '')}（只追加，不可篡改）<br/>"
        f"采集快照目录：{meta.get('snapshot_dir', '')}", st_body))
    story.append(Spacer(1, 4))
    story.append(Paragraph("说明：本报告全部指标由程序按配置限值计算生成，未经过模型推算；"
                           "每个数字均可通过上方指纹与记录号回溯至原始系统记录。"
                           "本报告仅作过程监控参考，放行决策以人工复核与检验报告为准。", st_sub))

    # ---------- 页面装饰（页眉线 + 页脚指纹/页码） ----------
    fp_short = meta.get("fingerprint", "")[:16]

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(18 * mm, 285 * mm, f"{meta.get('title', '食用油中控品控日报')} · {m.report_date}")
        canvas.drawRightString(192 * mm, 285 * mm, f"RunID {meta.get('run_id', '')}")
        canvas.setStrokeColor(RED)
        canvas.setLineWidth(1)
        canvas.line(18 * mm, 283 * mm, 192 * mm, 283 * mm)
        canvas.line(18 * mm, 14 * mm, 192 * mm, 14 * mm)
        canvas.drawString(18 * mm, 10 * mm, f"报告指纹 FP {fp_short}")
        canvas.drawRightString(192 * mm, 10 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=22 * mm, bottomMargin=18 * mm,
                            title=f"{meta.get('title', '品控日报')} {m.report_date}",
                            author=m.plant)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return path
